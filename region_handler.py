
import itertools
import openpyxl
import warnings
import re
import bed
import math

import dataclasses
import collections

from typing import *
from typing.io import *

"""
region_extractor = Region.Extractor(..., ...)
regions = region_extractor.extract_regions()
"""

REGION_COUNT = 4

TABLE_HEADING = [
    "Sequence",
    "Region Count",
    "#A",
    "#T",
    "#C",
    "#G",
    "Gene Count",
    "Weight",
]

warnings.simplefilter("ignore")


@dataclasses.dataclass
class RegionEntry:
    ntuple: str
    region_count: int
    nucleotide_count: Dict[str, int]
    gene_count: int
    weight: float


@dataclasses.dataclass
class RegionSummary:
    region: int
    entries: List[RegionEntry]


def _divide_interval(sequence: str, index: int, width: int, padding: int):
    first_interval = sequence[index - width - padding : index]
    second_interval = sequence[index : index + width + padding]

    # create the sliding windows
    first_regions = [
        first_interval[i : i + width] for i in range(len(first_interval) - width + 1)
    ]
    second_regions = [
        second_interval[i : i + width] for i in range(len(second_interval) - width + 1)
    ]

    return first_regions, second_regions


class Extractor:
    def __init__(self, genomic_regions: List[bed.GenomicRegion], sequence: str):
        self.genomic_regions = genomic_regions
        self.sequence = sequence

    def extract_all_regions(self, width, padding):
        for genomic_region in self.genomic_regions:
            yield self.extract_regions(genomic_region, width, padding)

    def extract_regions(
        self, genomic_region: bed.GenomicRegion, width: int, padding: int
    ):
        region1, region2 = _divide_interval(
            self.sequence, genomic_region.start, width, padding
        )
        region3, region4 = _divide_interval(
            self.sequence, genomic_region.end, width, padding
        )

        return (region1, region2, region3, region4)


def _increment_occurences(occurences_dict, regions):
    for window in regions:
        occurences_dict[window] += 1


def threshold_region_summaries_shannon(region_summaries: List[RegionSummary]):
    thresholded_region_summaries = [RegionSummary(i, []) for i in range(0, len(region_summaries))]
    for i, region_summary in enumerate(region_summaries):
        count = 1
        expected_surprise = 0
        max_weight = 0
        prev_entropy = -math.inf

        for entry in region_summary.entries:
            weight = entry.weight
            if count == 1 or max_weight == 0:
                max_weight = weight

            p = weight / max_weight
            surprise = -p * math.log(p, 10)
            expected_surprise += surprise
            entropy = expected_surprise / count

            if entropy >= prev_entropy:
                thresholded_region_summaries[i].entries.append(entry)
                prev_entropy = entropy
                count += 1

    return thresholded_region_summaries

def generate_region_summaries(
    sequence, rloops, width, padding, gene_start, gene_end
):
    regions_occurences = [
        collections.defaultdict(lambda: 0) for _ in range(REGION_COUNT)
    ]

    rloop_count = 0
    extractor = Extractor(rloops, sequence)
    for regions in extractor.extract_all_regions(width, padding):
        rloop_count += 1

        for i in range(len(regions)):
            _increment_occurences(regions_occurences[i], regions[i])

    region_summaries = [RegionSummary(i, []) for i in range(0, REGION_COUNT)]
    for i, region_occurences in enumerate(regions_occurences):
        for ntuple, region_count in region_occurences.items():
            a_count = ntuple.count("A")
            t_count = ntuple.count("T")
            c_count = ntuple.count("C")
            g_count = ntuple.count("G")

            nucleotide_count = {"A": a_count, "T": t_count, "C": c_count, "G": g_count}

            gene = sequence[gene_start:gene_end]
            gene_count = len(re.findall(f"(?={ntuple})", gene))
            weight = region_count / (gene_count * rloop_count)

            entry = RegionEntry(
                ntuple, region_count, nucleotide_count, gene_count, weight
            )
            region_summaries[i].entries.append(entry)

    return region_summaries


def write_region_summaries(region_summaries: List[RegionSummary], xlsx_filename):
    wb = openpyxl.Workbook(write_only=True)

    # build the xlsx document
    for i, region_summary in enumerate(region_summaries):
        ws = wb.create_sheet(f"Region {i + 1}", i)
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 15

        ws.column_dimensions["C"].width = 8
        ws.column_dimensions["D"].width = 8
        ws.column_dimensions["E"].width = 8
        ws.column_dimensions["F"].width = 8

        ws.column_dimensions["G"].width = 15
        ws.column_dimensions["H"].width = 15

        table = openpyxl.worksheet.table.Table(
            displayName=f"Region{i + 1}",
            ref=f"A1:H{len(region_summary.entries) + 1}",
        )

        table._initialise_columns()
        for column, value in zip(table.tableColumns, TABLE_HEADING):
            column.name = value

        style = openpyxl.worksheet.table.TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=True,
        )
        table.tableStyleInfo = style

        ws.add_table(table)

        ws.append(TABLE_HEADING)
        for entry in region_summary.entries:
            ws.append(
                [
                    entry.ntuple,
                    entry.region_count,
                    entry.nucleotide_count["A"],
                    entry.nucleotide_count["T"],
                    entry.nucleotide_count["C"],
                    entry.nucleotide_count["G"],
                    entry.gene_count,
                    entry.weight,
                ]
            )

    wb.save(xlsx_filename)


if __name__ == "__main__":
    region_summaries = generate_region_summaries(FASTA_FILE, BED_FILE, WIDTH, PADDING)
    write_region_summaries(region_summaries, "output.xlsx")
